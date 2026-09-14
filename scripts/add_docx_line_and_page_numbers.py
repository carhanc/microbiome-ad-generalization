#!/usr/bin/env python3
"""Adds Word's native continuous line numbering (section property) and a
page-number field in the footer to an already-built manuscript DOCX.

Uses only Word's own supported features (w:lnNumType, PAGE field) via direct
OOXML manipulation -- no content, table, or figure in the document is
touched or reflowed differently than pandoc already produced.
"""

import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX_PATH = "/Users/arhan/Desktop/microbiome-ad-generalization/manuscript/draft.docx"


def add_continuous_line_numbers(document):
    for section in document.sections:
        sectPr = section._sectPr
        # Remove any pre-existing lnNumType before adding ours.
        for existing in sectPr.findall(qn("w:lnNumType")):
            sectPr.remove(existing)
        lnNumType = OxmlElement("w:lnNumType")
        lnNumType.set(qn("w:countBy"), "1")
        lnNumType.set(qn("w:start"), "1")
        lnNumType.set(qn("w:restart"), "continuous")
        lnNumType.set(qn("w:distance"), "360")  # twips from text, ~0.25in
        sectPr.append(lnNumType)


def add_page_number_footer(document):
    for section in document.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        # Clear any existing paragraphs' runs, reuse the first paragraph.
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for run in list(p.runs):
            run._element.getparent().remove(run._element)
        p.alignment = 1  # center

        # Footer paragraphs are their own line-numbering "story"; without this,
        # some renderers (e.g. LibreOffice) number the footer paragraph itself
        # as "line 1" at the same left-margin x-position as the body's line
        # numbers, which reads as a second, spurious number next to the true
        # centered PAGE field.
        pPr = p._p.get_or_add_pPr()
        for existing in pPr.findall(qn("w:suppressLineNumbers")):
            pPr.remove(existing)
        suppress = OxmlElement("w:suppressLineNumbers")
        pPr.insert(0, suppress)

        run = p.add_run()
        fldChar_begin = OxmlElement("w:fldChar")
        fldChar_begin.set(qn("w:fldCharType"), "begin")
        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = "PAGE"
        fldChar_sep = OxmlElement("w:fldChar")
        fldChar_sep.set(qn("w:fldCharType"), "separate")
        fldChar_end = OxmlElement("w:fldChar")
        fldChar_end.set(qn("w:fldCharType"), "end")

        r_element = run._element
        r_element.append(fldChar_begin)
        r_element.append(instrText)
        r_element.append(fldChar_sep)
        r_element.append(fldChar_end)


def main():
    doc = Document(DOCX_PATH)
    add_continuous_line_numbers(doc)
    add_page_number_footer(doc)
    doc.save(DOCX_PATH)
    print(f"Added continuous line numbering (section property) and PAGE-field footer to {DOCX_PATH}")


if __name__ == "__main__":
    main()
